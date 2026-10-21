"""
AI Interviewer - local app.
Run:  python app.py   then open http://127.0.0.1:5000
All data (settings, recordings, evaluations, reports) is stored next to this file.
"""
import base64
import csv
import io
import json
import os
import re
import threading
import time
import uuid
import webbrowser
from datetime import datetime

from flask import Flask, jsonify, request, send_from_directory, abort

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
SESSIONS_DIR = os.path.join(BASE_DIR, "sessions")
QUESTION_SETS_DIR = os.path.join(DATA_DIR, "question_sets")
SETTINGS_FILE = os.path.join(DATA_DIR, "settings.json")
for d in (DATA_DIR, SESSIONS_DIR, QUESTION_SETS_DIR):
    os.makedirs(d, exist_ok=True)

app = Flask(__name__, static_folder="static", static_url_path="/static")
app.config["MAX_CONTENT_LENGTH"] = 1024 * 1024 * 1024  # 1 GB uploads (video)

LOCK = threading.Lock()

DEFAULT_SETTINGS = {
    "openai_api_key": "",
    "chat_model": "gpt-6-luna",
    "transcribe_model": "gpt-4o-mini-transcribe",
    "analyze_body_language": True,
    "speak_questions": True,
    "candidate_name": "",
    "target_role": "Software Engineer",
    "experience_years": "5",
    "tech_stack": "",
    "job_description": "",
    "answer_time_limit_sec": 180,
    "tts_model": "gpt-4o-mini-tts",
    "audio_model": "gpt-audio-mini",
    "llm_provider": "openai",
    "llm_api_key": "",
    "llm_base_url": "",
}


# ----------------------------------------------------------------- helpers
def read_json(path, default=None):
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return default


def write_json(path, obj):
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(obj, f, indent=2, ensure_ascii=False)
    os.replace(tmp, path)


def load_settings():
    s = dict(DEFAULT_SETTINGS)
    s.update(read_json(SETTINGS_FILE, {}) or {})
    return s


def save_settings(s):
    write_json(SETTINGS_FILE, s)


def mask_key(k):
    if not k:
        return ""
    return k[:5] + "..." + k[-4:] if len(k) > 12 else "****"


# ----------------------------------------------------------- AI providers
# Scoring/coaching works with any OpenAI-compatible chat API. Voice features (speech-to-text,
# natural read-aloud voice, "AI listens") use OpenAI speech models and need an OpenAI key.
PROVIDERS = {
    "openai": "",
    "anthropic": "https://api.anthropic.com/v1/",
    "gemini": "https://generativelanguage.googleapis.com/v1beta/openai/",
    "groq": "https://api.groq.com/openai/v1",
    "openrouter": "https://openrouter.ai/api/v1",
    "ollama": "http://localhost:11434/v1",
    "custom": "",
}
VOICE_HELP = ("This feature uses OpenAI speech models, so it needs an OpenAI API key. Add one in Settings "
              "(you can keep another provider for scoring), or type your answer / switch to Browser voice.")


def client():
    """OpenAI client for voice features (speech-to-text, text-to-speech, AI listens)."""
    from openai import OpenAI
    key = load_settings().get("openai_api_key")
    if not key:
        raise RuntimeError(VOICE_HELP)
    return OpenAI(api_key=key)


def llm_config(s=None):
    """(provider, base_url, api_key) for the scoring/coaching model."""
    s = s or load_settings()
    p = s.get("llm_provider") or "openai"
    if p not in PROVIDERS:
        p = "custom"
    if p == "openai":
        return p, None, s.get("openai_api_key") or ""
    base = (s.get("llm_base_url") or "").strip() if p == "custom" else PROVIDERS[p]
    key = s.get("llm_api_key") or ("ollama" if p == "ollama" else "")
    return p, base, key


def llm_ready(s):
    p, base, key = llm_config(s)
    return bool(key) and (p != "custom" or bool(base))


def chat_client(s=None):
    from openai import OpenAI
    p, base, key = llm_config(s)
    if p == "custom" and not base:
        raise RuntimeError("The custom AI provider needs a base URL. Open Settings and add it.")
    if not key:
        raise RuntimeError("No API key is set for the AI provider. Open Settings and add it.")
    return OpenAI(api_key=key, base_url=base) if base else OpenAI(api_key=key)


def stt():
    """(client, model) for speech-to-text: OpenAI when a key is set, else Groq's Whisper when Groq is the provider."""
    s = load_settings()
    if s.get("openai_api_key"):
        return client(), s["transcribe_model"]
    if s.get("llm_provider") == "groq" and s.get("llm_api_key"):
        return chat_client(s), "whisper-large-v3-turbo"
    raise RuntimeError(VOICE_HELP)


def _has_images(messages):
    return any(isinstance(m.get("content"), list) and any(x.get("type") == "image_url" for x in m["content"])
               for m in messages)


def _flatten(messages):
    """Turn list content into plain text (drops images) for providers that only accept strings."""
    out = []
    for m in messages:
        c = m.get("content")
        if isinstance(c, list):
            c = "\n".join(x.get("text", "") for x in c if x.get("type") == "text")
        out.append({**m, "content": c})
    return out


def chat_create(messages, max_tokens, json_mode=False, s=None):
    """Chat call that works across OpenAI-compatible providers, dropping options a provider rejects."""
    s = s or load_settings()
    p = llm_config(s)[0]
    if p != "openai" and not _has_images(messages):
        messages = _flatten(messages)
    kw = {"model": s["chat_model"], "messages": messages}
    kw["max_completion_tokens" if p == "openai" else "max_tokens"] = max_tokens
    if json_mode:
        kw["response_format"] = {"type": "json_object"}
    oc = chat_client(s)
    for _ in range(3):
        try:
            return oc.chat.completions.create(**kw)
        except Exception as e:
            msg = str(e).lower()
            if "response_format" in kw and ("response_format" in msg or "json" in msg):
                kw.pop("response_format")
                continue
            if _has_images(kw["messages"]) and any(w in msg for w in ("image", "vision", "multimodal", "content")):
                kw["messages"] = _flatten(kw["messages"])  # model can't see images: score without them
                continue
            raise
    return oc.chat.completions.create(**kw)


def profile_text(s, jd=None):
    """jd: job description for this interview (overrides the one in Settings when given)."""
    jd = (jd or "").strip() or (s.get("job_description") or "").strip()
    lines = [
        f"Target role: {s.get('target_role') or 'not specified'}",
        f"Years of experience the candidate wants to sound like: {s.get('experience_years') or 'not specified'}",
    ]
    if s.get("tech_stack"):
        lines.append(f"Tech stack / domain: {s['tech_stack']}")
    if jd:
        lines.append("Job description (tailor expectations, vocabulary and examples to this):\n" + jd[:6000])
    return "\n".join(lines)


def chat_json(system, user_content, max_tokens=4000):
    resp = chat_create([{"role": "system", "content": system}, {"role": "user", "content": user_content}],
                       max_tokens, json_mode=True)
    txt = resp.choices[0].message.content or "{}"
    try:
        return json.loads(txt)
    except json.JSONDecodeError:
        m = re.search(r"\{.*\}", txt, re.S)
        return json.loads(m.group(0)) if m else {"error": "Model did not return JSON", "raw": txt}


def chat_text(system, user_content, max_tokens=4000):
    resp = chat_create([{"role": "system", "content": system}, {"role": "user", "content": user_content}], max_tokens)
    return resp.choices[0].message.content or ""


def slug(t, n=40):
    t = re.sub(r"[^A-Za-z0-9]+", "_", t or "").strip("_")
    return t[:n] or "session"


# ------------------------------------------------------- question parsing
Q_START = re.compile(r"^\s*(?:Q(?:uestion)?\s*\d*[\.\):\-]?|\d{1,3}\s*[\.\)\:\-]|[-*•]\s+)\s*", re.I)


def split_questions_from_text(text):
    """Heuristic: numbered/bulleted items or lines ending with '?' become questions."""
    lines = [l.strip() for l in text.replace("\r", "").split("\n")]
    questions, current = [], ""
    numbered = sum(1 for l in lines if Q_START.match(l) and len(Q_START.sub("", l)) > 3)
    if numbered >= 2:
        for l in lines:
            if not l:
                continue
            if Q_START.match(l) and len(Q_START.sub("", l)) > 3:
                if current:
                    questions.append(current.strip())
                current = Q_START.sub("", l)
            elif current and not re.match(r"^(answer|ans|a)\s*[:\-]", l, re.I):
                # continuation line of a question (only if question not finished)
                if not current.rstrip().endswith("?"):
                    current += " " + l
        if current:
            questions.append(current.strip())
    else:
        for l in lines:
            if l.endswith("?") and len(l) > 5:
                questions.append(l)
        if not questions:
            questions = [l for l in lines if len(l) > 8]
    # de-dup
    seen, out = set(), []
    for q in questions:
        k = q.lower()
        if k not in seen:
            seen.add(k)
            out.append({"question": q, "reference_answer": "", "category": ""})
    return out


def parse_excel(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = [r for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
        if not rows:
            continue
        out.extend(rows_to_questions(rows))
    return out


def rows_to_questions(rows):
    header = [str(c or "").strip().lower() for c in rows[0]]
    qi = next((i for i, h in enumerate(header) if "question" in h), None)
    ai = next((i for i, h in enumerate(header) if any(w in h for w in ("answer", "expected", "solution"))), None)
    ci = next((i for i, h in enumerate(header) if any(w in h for w in ("category", "topic", "type", "section"))), None)
    body = rows[1:] if qi is not None else rows
    if qi is None:
        # pick the column with the longest average text
        ncols = max(len(r) for r in rows)
        lens = [sum(len(str(r[i] or "")) for r in rows if i < len(r)) for i in range(ncols)]
        qi = lens.index(max(lens))
    out = []
    for r in body:
        q = str(r[qi] or "").strip() if qi < len(r) else ""
        if len(q) < 3:
            continue
        out.append({
            "question": q,
            "reference_answer": str(r[ai] or "").strip() if ai is not None and ai < len(r) else "",
            "category": str(r[ci] or "").strip() if ci is not None and ci < len(r) else "",
        })
    return out


def parse_csv(data):
    text = data.decode("utf-8-sig", errors="ignore")
    rows = [r for r in csv.reader(io.StringIO(text)) if any(c.strip() for c in r)]
    return rows_to_questions(rows) if rows else []


def parse_pdf(data):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return "\n".join((p.extract_text() or "") for p in reader.pages)


def parse_docx(data):
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def ai_extract_questions(text):
    sys = ("You extract interview questions from documents. Return JSON: "
           '{"questions":[{"question":"...","reference_answer":"(answer if the document provides one, else empty)",'
           '"category":"(e.g. Technical, Behavioral, System Design, HR)"}]}. '
           "Keep each question's wording faithful. Do not invent questions that are not in the text.")
    res = chat_json(sys, text[:60000], max_tokens=8000)
    return [q for q in res.get("questions", []) if q.get("question")]


# --------------------------------------------------------------- metrics
FILLERS = ["um", "uh", "erm", "ah", "like", "you know", "basically", "actually", "literally",
           "sort of", "kind of", "i mean", "so yeah", "right", "okay so", "stuff", "things like that"]


def speech_metrics(transcript, duration_sec):
    words = re.findall(r"[A-Za-z']+", transcript or "")
    low = " " + " ".join(w.lower() for w in words) + " "
    fillers = {}
    for f in FILLERS:
        c = low.count(" " + f + " ")
        if c:
            fillers[f] = c
    dur = float(duration_sec or 0)
    return {
        "word_count": len(words),
        "duration_sec": round(dur, 1),
        "words_per_minute": round(len(words) / (dur / 60.0)) if (words and dur >= 1) else 0,
        "filler_words": fillers,
        "filler_total": sum(fillers.values()),
    }


# ------------------------------------------------------------ AI prompts
EVAL_SYSTEM = """You are a senior hiring manager and an elite interview coach. You evaluate one spoken interview answer
and then coach the candidate so they sound like a genuinely experienced professional in the target role.
Be honest and specific - no generic advice. Quote the candidate's own words when pointing out problems.
Scores are 0-10 (10 = would impress a bar-raiser at a top company; 5 = borderline; below 4 = would fail).

Return ONLY JSON with exactly these keys:
{
 "score": number 0-10 (one decimal),
 "verdict": "Strong Hire | Hire | Lean Hire | Lean No Hire | No Hire",
 "summary": "2-3 sentence overall judgement",
 "scores": {"technical_accuracy":n, "depth":n, "structure":n, "communication":n, "confidence":n, "relevance":n},
 "what_you_did_well": ["..."],
 "improvements": ["specific, actionable things to fix, most important first"],
 "missing_points": ["key points/concepts an experienced person would have mentioned but candidate did not"],
 "mistakes": ["factually wrong or misleading statements, with the correction"],
 "technical_terms_to_use": [{"term":"...", "how_to_use":"a sentence showing it used naturally in this answer"}],
 "better_phrasing": [{"you_said":"exact or close quote", "say_instead":"how a senior person would phrase it"}],
 "answer_structure": "the best framework for this question (e.g. STAR, Definition->How it works->Example->Trade-offs->When to use) and how to apply it here",
 "model_answer": "a complete answer written exactly as an experienced professional would SAY it out loud in an interview: first person, natural spoken tone, includes a short real-world experience example ('In my last project we...'), correct terminology, trade-offs, and a crisp close. 150-300 words.",
 "delivery_feedback": {"pace":"...", "filler_words":"...", "clarity_and_conciseness":"...", "confidence_and_tone":"...", "body_language":"(only if frames provided, else 'Not analyzed')"},
 "likely_follow_ups": [{"question":"...", "how_to_answer":"short guidance"}],
 "practice_drill": "one concrete exercise to do before re-attempting this question",
 "tone_authenticity": {
   "rating": "Natural & genuine | Mostly natural | Somewhat rehearsed | Sounds memorized / read from notes",
   "score": number 0-10 (10 = completely natural, conversational, clearly from own experience),
   "verdict": "2-3 plain sentences telling the candidate honestly how they come across to an interviewer",
   "evidence": ["specific signals you noticed, quoting their words: textbook/definition-style sentences, no personal examples, list-like recitation, unnaturally perfect flow with no thinking pauses, flat steady volume, eyes looking away/down at notes in frames - OR the positive signals: own project stories, 'we/I' ownership, natural pauses, opinions, trade-offs from experience"],
   "how_to_sound_natural": ["concrete tips to sound genuine and experienced instead of studied"]
 },
 "progress_vs_previous": null OR (only when a previous attempt is provided) {
   "previous_score": number, "score_change": number (this score minus previous),
   "overall": "1-2 sentences: did they get better, worse or stay the same, and why",
   "improved": ["things that are clearly better than last attempt, quoting both versions where useful"],
   "not_improved": ["feedback from last time that they still did not fix"],
   "new_issues": ["problems that appeared this time but were not there before"],
   "next_focus": "the single most important thing to fix on the next try"
 }
}

How to judge tone_authenticity: an answer that sounds read or memorized usually has written-language sentences
("X is a mechanism which facilitates..."), recites definitions in textbook order, lacks first-person examples,
has almost no filler words or thinking pauses, and steady volume. A genuine answer mixes explanation with personal
experience, reasoning out loud, opinions, and natural pauses. Use the transcript, the speech and voice metrics, and
the webcam frames (eyes repeatedly off-camera or down = likely reading). Voice metrics are approximate.
Judge the previous-attempt comparison strictly against the previous feedback - do not invent progress."""

LEARN_SYSTEM = """You are a world-class interview coach and senior practitioner. Teach the candidate how to answer the given
interview question from absolute basics to expert level so they can answer like someone who has years of real experience.
Write in clear Markdown with these sections:
## 1. What the interviewer is really testing
## 2. Concept from zero (simple explanation, analogy)
## 3. Going deeper (how it works, internals, edge cases, trade-offs)
## 4. Key technical terms & vocabulary (table: Term | Meaning | Use it like this)
## 5. Best answer structure (framework, with time budget for a 2-minute answer)
## 6. Model answer (spoken, first person, as an experienced professional, include a realistic project story)
## 7. How to sound experienced (phrases seniors use, phrases that sound junior and to avoid)
## 8. Common mistakes
## 9. Likely follow-up questions with short answers
## 10. 5-minute practice drill
Tailor everything to the candidate profile."""

REPORT_SYSTEM = """You are an elite interview coach writing a complete training report after a mock interview.
Use the per-question evaluations provided. Be specific, honest and motivating. Return ONLY JSON:
{
 "overall_score": number 0-10,
 "readiness": "Ready | Almost ready | Needs practice | Not ready",
 "headline": "one sentence",
 "strengths": ["..."],
 "weaknesses": ["..."],
 "communication_coaching": ["how to speak like an experienced professional: pace, structure, tone, storytelling, confidence"],
 "vocabulary_to_master": [{"term":"...", "meaning":"...", "example":"..."}],
 "topics_to_study": [{"topic":"...", "why":"...", "what_to_cover":"..."}],
 "seven_day_plan": [{"day":1, "focus":"...", "tasks":["..."]}],
 "phrases_bank": ["ready-to-use senior-sounding phrases for openings, transitions, trade-offs, closings"],
 "tone_summary": "how natural/genuine vs rehearsed the candidate sounds overall and the #1 fix",
 "next_session_advice": "..."
}"""


# --------------------------------------------------------------- routes: UI
@app.route("/")
def index():
    return send_from_directory(os.path.join(BASE_DIR, "static"), "index.html")


@app.route("/media/<session_id>/<path:fname>")
def media(session_id, fname):
    folder = session_folder(session_id)
    if not folder:
        abort(404)
    return send_from_directory(folder, fname)


# --------------------------------------------------------- routes: settings
@app.get("/api/settings")
def get_settings():
    s = load_settings()
    out = dict(s)
    out["has_key"] = llm_ready(s)  # scoring/coaching provider is configured
    out["has_openai_key"] = bool(s["openai_api_key"])
    out["has_llm_key"] = bool(s.get("llm_api_key"))
    out["openai_api_key"] = mask_key(s["openai_api_key"])
    out["llm_api_key"] = mask_key(s.get("llm_api_key"))
    return jsonify(out)


@app.post("/api/settings")
def post_settings():
    body = request.get_json(force=True) or {}
    s = load_settings()
    secrets = ("openai_api_key", "llm_api_key")
    for k in DEFAULT_SETTINGS:
        if k in body and k not in secrets:
            s[k] = body[k]
    for k in secrets:
        newkey = (body.get(k) or "").strip()
        if newkey and "..." not in newkey:
            s[k] = newkey
        elif body.get(k) is None and body.get("clear_" + k):
            s[k] = ""
    save_settings(s)
    return get_settings()


@app.post("/api/settings/test")
def test_key():
    """kind=openai: check the OpenAI key (voice features). Otherwise: send a tiny chat to the AI provider."""
    body = request.get_json(force=True) or {}
    s = load_settings()

    def pick(name):
        v = (body.get(name) or "").strip()
        return v if v and "..." not in v else s.get(name, "")
    try:
        from openai import OpenAI
        if body.get("kind") == "openai":
            OpenAI(api_key=pick("openai_api_key")).models.list()
        else:
            t = dict(s)
            for k in ("llm_provider", "llm_base_url", "chat_model"):
                if body.get(k):
                    t[k] = body[k]
            t["llm_api_key"], t["openai_api_key"] = pick("llm_api_key"), pick("openai_api_key")
            chat_create([{"role": "user", "content": "Reply with the word OK."}], 50, s=t)
        return jsonify({"ok": True})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 400


# -------------------------------------------------------- routes: questions
@app.post("/api/questions/parse")
def parse_questions():
    """Accepts a file (xlsx, csv, pdf, docx, txt) or JSON {text}. Returns questions."""
    use_ai = request.args.get("ai") == "1"
    text = None
    questions = []
    if "file" in request.files:
        f = request.files["file"]
        name = (f.filename or "").lower()
        data = f.read()
        try:
            if name.endswith((".xlsx", ".xlsm")):
                questions = parse_excel(data)
            elif name.endswith(".csv"):
                questions = parse_csv(data)
            elif name.endswith(".pdf"):
                text = parse_pdf(data)
            elif name.endswith(".docx"):
                text = parse_docx(data)
            elif name.endswith((".txt", ".md")):
                text = data.decode("utf-8", errors="ignore")
            else:
                return jsonify({"error": "Unsupported file type. Use .xlsx, .csv, .pdf, .docx or .txt"}), 400
        except Exception as e:
            return jsonify({"error": f"Could not read file: {e}"}), 400
    else:
        body = request.get_json(force=True, silent=True) or {}
        text = body.get("text", "")
    if text is not None:
        if use_ai:
            try:
                questions = ai_extract_questions(text)
            except Exception as e:
                return jsonify({"error": f"AI extraction failed: {e}"}), 400
        else:
            questions = split_questions_from_text(text)
    return jsonify({"questions": questions})


@app.get("/api/question-sets")
def list_sets():
    out = []
    for fn in sorted(os.listdir(QUESTION_SETS_DIR)):
        if fn.endswith(".json"):
            d = read_json(os.path.join(QUESTION_SETS_DIR, fn), {})
            out.append({"id": fn[:-5], "name": d.get("name"), "count": len(d.get("questions", []))})
    return jsonify(out)


@app.get("/api/question-sets/<sid>")
def get_set(sid):
    p = os.path.join(QUESTION_SETS_DIR, slug(sid, 80) + ".json")
    d = read_json(p)
    return jsonify(d) if d else (jsonify({"error": "not found"}), 404)


@app.post("/api/question-sets")
def save_set():
    body = request.get_json(force=True) or {}
    name = body.get("name") or "Question set"
    sid = slug(name, 60)
    write_json(os.path.join(QUESTION_SETS_DIR, sid + ".json"),
               {"name": name, "questions": body.get("questions", [])})
    return jsonify({"id": sid})


@app.delete("/api/question-sets/<sid>")
def delete_set(sid):
    p = os.path.join(QUESTION_SETS_DIR, slug(sid, 80) + ".json")
    if os.path.exists(p):
        os.remove(p)
    return jsonify({"ok": True})


@app.post("/api/learn")
def learn():
    body = request.get_json(force=True) or {}
    q = body.get("question", "")
    ref = body.get("reference_answer", "")
    s = load_settings()
    user = f"Candidate profile:\n{profile_text(s, body.get('job_description'))}\n\nInterview question:\n{q}"
    if ref:
        user += f"\n\nReference answer provided by the candidate's material:\n{ref}"
    try:
        return jsonify({"markdown": chat_text(LEARN_SYSTEM, user, max_tokens=5000)})
    except Exception as e:
        return jsonify({"error": str(e)}), 400


# --------------------------------------------------------- routes: sessions
def session_folder(session_id):
    sid = slug(session_id, 120)
    p = os.path.join(SESSIONS_DIR, sid)
    return p if os.path.isdir(p) else None


def load_session(session_id):
    f = session_folder(session_id)
    return read_json(os.path.join(f, "session.json")) if f else None


def save_session(sess):
    write_json(os.path.join(SESSIONS_DIR, sess["id"], "session.json"), sess)


@app.post("/api/sessions")
def create_session():
    body = request.get_json(force=True) or {}
    qs = [q for q in body.get("questions", []) if (q.get("question") or "").strip()]
    if not qs:
        return jsonify({"error": "No questions"}), 400
    name = body.get("name") or "Interview"
    sid = datetime.now().strftime("%Y%m%d_%H%M%S") + "_" + slug(name, 30)
    os.makedirs(os.path.join(SESSIONS_DIR, sid), exist_ok=True)
    s = load_settings()
    sess = {
        "id": sid, "name": name, "created": datetime.now().isoformat(timespec="seconds"),
        "mode": body.get("mode", "coach"),
        "job_description": (body.get("job_description") or "").strip(),
        "profile": {k: s.get(k) for k in ("target_role", "experience_years", "tech_stack", "job_description")},
        "questions": [{"index": i, "question": q["question"].strip(),
                       "reference_answer": q.get("reference_answer", ""),
                       "category": q.get("category", ""), "attempts": []} for i, q in enumerate(qs)],
        "report": None,
    }
    save_session(sess)
    return jsonify(sess)


@app.get("/api/sessions")
def list_sessions():
    out = []
    for sid in sorted(os.listdir(SESSIONS_DIR), reverse=True):
        sess = read_json(os.path.join(SESSIONS_DIR, sid, "session.json"))
        if not sess:
            continue
        scores = []
        for q in sess["questions"]:
            done = [a for a in q["attempts"] if a.get("status") == "done" and a.get("evaluation")]
            if done:
                scores.append(done[-1]["evaluation"].get("score", 0) or 0)
        out.append({"id": sid, "name": sess["name"], "created": sess["created"],
                    "questions": len(sess["questions"]), "answered": len(scores),
                    "avg_score": round(sum(scores) / len(scores), 1) if scores else None,
                    "has_report": bool(sess.get("report"))})
    return jsonify(out)


@app.get("/api/sessions/<sid>")
def get_session(sid):
    sess = load_session(sid)
    return jsonify(sess) if sess else (jsonify({"error": "not found"}), 404)


@app.delete("/api/sessions/<sid>")
def delete_session(sid):
    import shutil
    f = session_folder(sid)
    if f:
        shutil.rmtree(f, ignore_errors=True)
    return jsonify({"ok": True})


@app.post("/api/sessions/<sid>/questions/<int:qi>/answer")
def submit_answer(sid, qi):
    """multipart: video (webm), audio (webm), duration, frames (json list of dataURLs)."""
    with LOCK:
        sess = load_session(sid)
        if not sess or qi >= len(sess["questions"]):
            return jsonify({"error": "not found"}), 404
        q = sess["questions"][qi]
        n = len(q["attempts"]) + 1
        folder = os.path.join(SESSIONS_DIR, sess["id"])
        base = f"q{qi + 1:02d}_attempt{n}"
        att = {"attempt": n, "created": datetime.now().isoformat(timespec="seconds"),
               "status": "processing", "stage": "saving", "video": None, "audio": None, "frames": [],
               "duration_sec": float(request.form.get("duration") or 0)}
        v = request.files.get("video")
        a = request.files.get("audio")
        if v:
            ext = ".mp4" if "mp4" in (v.mimetype or "") else ".webm"
            att["video"] = base + "_video" + ext
            v.save(os.path.join(folder, att["video"]))
        if a:
            ext = ".mp4" if "mp4" in (a.mimetype or "") else (".ogg" if "ogg" in (a.mimetype or "") else ".webm")
            att["audio"] = base + "_audio" + ext
            a.save(os.path.join(folder, att["audio"]))
        try:
            vm = json.loads(request.form.get("voice_metrics") or "null")
            if isinstance(vm, dict):
                att["voice_metrics"] = vm
        except Exception:
            pass
        typed = (request.form.get("typed_answer") or "").strip()
        if typed:
            att["typed_answer"] = typed
        try:
            frames = json.loads(request.form.get("frames") or "[]")
        except Exception:
            frames = []
        for i, fr in enumerate(frames[:4]):
            m = re.match(r"data:image/\w+;base64,(.*)", fr or "")
            if m:
                fn = f"{base}_frame{i + 1}.jpg"
                with open(os.path.join(folder, fn), "wb") as fh:
                    fh.write(base64.b64decode(m.group(1)))
                att["frames"].append(fn)
        q["attempts"].append(att)
        save_session(sess)
    threading.Thread(target=process_attempt, args=(sess["id"], qi, n - 1), daemon=True).start()
    return jsonify({"ok": True, "attempt": n})


def update_attempt(sid, qi, ai, **kw):
    with LOCK:
        sess = load_session(sid)
        sess["questions"][qi]["attempts"][ai].update(kw)
        save_session(sess)
        return sess


def find_previous_attempt(sess, q, ai):
    """Latest evaluated attempt of the same question: first in this session, else in older sessions."""
    same = [a for a in q["attempts"][:ai] if a.get("status") == "done" and a.get("evaluation")]
    if same:
        return same[-1]
    key = q["question"].strip().lower()
    best = None
    for sid in os.listdir(SESSIONS_DIR):
        if sid == sess["id"]:
            continue
        other = read_json(os.path.join(SESSIONS_DIR, sid, "session.json"))
        for oq in (other or {}).get("questions", []):
            if oq["question"].strip().lower() != key:
                continue
            for a in oq["attempts"]:
                if a.get("status") == "done" and a.get("evaluation") and (not best or a["created"] > best["created"]):
                    best = a
    return best


def process_attempt(sid, qi, ai):
    try:
        sess = load_session(sid)
        q = sess["questions"][qi]
        att = q["attempts"][ai]
        folder = os.path.join(SESSIONS_DIR, sid)
        s = load_settings()

        # 1) transcription
        transcript = att.get("typed_answer", "")
        if att.get("audio") and not transcript:
            update_attempt(sid, qi, ai, stage="transcribing")
            path = os.path.join(folder, att["audio"])
            if os.path.getsize(path) > 24 * 1024 * 1024:
                raise RuntimeError("Audio file is larger than 25 MB, which is the OpenAI transcription limit. Keep answers shorter.")
            oc, stt_model = stt()
            with open(path, "rb") as fh:
                tr = oc.audio.transcriptions.create(model=stt_model, file=fh,
                                                    prompt="Technical job interview answer. Keep filler words like um, uh, like.")
            transcript = getattr(tr, "text", "") or ""
        with open(os.path.join(folder, f"q{qi + 1:02d}_attempt{att['attempt']}_transcript.txt"), "w", encoding="utf-8") as fh:
            fh.write(transcript)
        metrics = speech_metrics(transcript, att.get("duration_sec"))
        update_attempt(sid, qi, ai, transcript=transcript, metrics=metrics, stage="evaluating")

        # 2) evaluation
        prev = find_previous_attempt(sess, q, ai)
        text = (f"Candidate profile:\n{profile_text(s, sess.get('job_description'))}\n\n"
                f"Question category: {q.get('category') or 'unknown'}\n"
                f"Interview question:\n{q['question']}\n\n")
        if q.get("reference_answer"):
            text += f"Reference/expected answer (from candidate's material):\n{q['reference_answer']}\n\n"
        text += (f"Candidate's spoken answer (auto-transcribed):\n\"\"\"{transcript or '(no speech detected)'}\"\"\"\n\n"
                 f"Speech metrics: {json.dumps(metrics)}\n"
                 "(Ideal interview pace is ~130-160 words per minute.)\n")
        if att.get("voice_metrics"):
            text += f"Voice metrics from the microphone (approximate): {json.dumps(att['voice_metrics'])}\n"
        if prev:
            pev = prev["evaluation"]
            text += ("\n=== PREVIOUS ATTEMPT AT THIS SAME QUESTION ===\n"
                     f"When: {prev.get('created')}\nPrevious score: {pev.get('score')}  sub-scores: {json.dumps(pev.get('scores'))}\n"
                     f"Previous tone rating: {(pev.get('tone_authenticity') or {}).get('rating')}\n"
                     f"Previous transcript:\n\"\"\"{(prev.get('transcript') or '')[:4000]}\"\"\"\n"
                     f"Feedback given last time - improvements: {json.dumps(pev.get('improvements'))}\n"
                     f"Missing points last time: {json.dumps(pev.get('missing_points'))}\n"
                     f"Delivery feedback last time: {json.dumps(pev.get('delivery_feedback'))}\n"
                     "Fill progress_vs_previous by comparing this attempt with that one.\n")
        else:
            text += "\nNo previous attempt: set progress_vs_previous to null.\n"
        content = [{"type": "text", "text": text}]
        if s.get("analyze_body_language") and att.get("frames"):
            content[0]["text"] += "\nAttached are webcam frames captured while answering. Assess eye contact, posture, expression, framing, lighting and professionalism in delivery_feedback.body_language.\n"
            for fn in att["frames"]:
                with open(os.path.join(folder, fn), "rb") as fh:
                    b64 = base64.b64encode(fh.read()).decode()
                content.append({"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{b64}", "detail": "low"}})
        ev = chat_json(EVAL_SYSTEM, content, max_tokens=6000)
        if prev and isinstance(ev.get("progress_vs_previous"), dict):
            try:
                ev["progress_vs_previous"]["previous_score"] = prev["evaluation"].get("score")
                ev["progress_vs_previous"]["score_change"] = round(float(ev.get("score") or 0) - float(prev["evaluation"].get("score") or 0), 1)
            except Exception:
                pass
        write_json(os.path.join(folder, f"q{qi + 1:02d}_attempt{att['attempt']}_evaluation.json"), ev)
        update_attempt(sid, qi, ai, evaluation=ev, status="done", stage="done")
    except Exception as e:
        update_attempt(sid, qi, ai, status="error", stage="error", error=str(e))


@app.post("/api/sessions/<sid>/report")
def make_report(sid):
    sess = load_session(sid)
    if not sess:
        return jsonify({"error": "not found"}), 404
    items = []
    for q in sess["questions"]:
        done = [a for a in q["attempts"] if a.get("status") == "done"]
        if not done:
            continue
        a = done[-1]
        ev = a["evaluation"]
        items.append({"question": q["question"], "score": ev.get("score"), "scores": ev.get("scores"),
                      "summary": ev.get("summary"), "improvements": ev.get("improvements"),
                      "missing_points": ev.get("missing_points"), "delivery": ev.get("delivery_feedback"), "tone": ev.get("tone_authenticity"),
                      "attempts": len(done), "progress": ev.get("progress_vs_previous"),
                      "metrics": a.get("metrics"), "transcript": (a.get("transcript") or "")[:1500]})
    if not items:
        return jsonify({"error": "No evaluated answers yet."}), 400
    s = load_settings()
    try:
        rep = chat_json(REPORT_SYSTEM, f"Candidate profile:\n{profile_text(s, sess.get('job_description'))}\n\nPer-question results:\n{json.dumps(items, indent=1)}", 6000)
    except Exception as e:
        return jsonify({"error": str(e)}), 400
    rep["generated"] = datetime.now().isoformat(timespec="seconds")
    with LOCK:
        sess = load_session(sid)
        sess["report"] = rep
        save_session(sess)
    write_markdown_report(sess)
    return jsonify(rep)


def write_markdown_report(sess):
    r = sess.get("report") or {}
    L = [f"# Interview Report - {sess['name']}", f"_{sess['created']}_", ""]
    if r:
        L += [f"**Overall score:** {r.get('overall_score')}/10 - **{r.get('readiness')}**", "", r.get("headline", ""), ""]
        for title, key in (("Strengths", "strengths"), ("Weaknesses", "weaknesses"),
                           ("Communication coaching", "communication_coaching"), ("Phrases bank", "phrases_bank")):
            L += [f"## {title}"] + [f"- {x}" for x in r.get(key, [])] + [""]
        L += ["## Vocabulary to master"] + [f"- **{v.get('term')}**: {v.get('meaning')} _e.g._ {v.get('example')}" for v in r.get("vocabulary_to_master", [])] + [""]
        L += ["## Topics to study"] + [f"- **{t.get('topic')}** - {t.get('why')} ({t.get('what_to_cover')})" for t in r.get("topics_to_study", [])] + [""]
        L += ["## 7-day plan"]
        for d in r.get("seven_day_plan", []):
            L += [f"### Day {d.get('day')}: {d.get('focus')}"] + [f"- {t}" for t in d.get("tasks", [])]
        if r.get("tone_summary"):
            L += ["", f"**Tone overall:** {r['tone_summary']}"]
        L += ["", f"**Next session:** {r.get('next_session_advice', '')}", ""]
    L.append("# Question by question")
    for q in sess["questions"]:
        done = [a for a in q["attempts"] if a.get("status") == "done"]
        L += ["", f"## Q{q['index'] + 1}. {q['question']}"]
        if not done:
            L.append("_Not answered._")
            continue
        a = done[-1]
        ev = a["evaluation"]
        L += [f"**Score:** {ev.get('score')}/10 ({ev.get('verdict')})", "", ev.get("summary", ""), "",
              "**Your answer (transcript):**", "", "> " + (a.get("transcript") or "").replace("\n", " "), "",
              "**Improvements:**"] + [f"- {x}" for x in ev.get("improvements", [])]
        t = ev.get("tone_authenticity") or {}
        if t:
            L += ["", f"**Tone:** {t.get('rating')} ({t.get('score')}/10) - {t.get('verdict', '')}"]
        pr = ev.get("progress_vs_previous") or {}
        if pr:
            L += ["", f"**Progress vs previous attempt:** {pr.get('previous_score')} -> {ev.get('score')} ({pr.get('score_change'):+}) - {pr.get('overall', '')}" if isinstance(pr.get('score_change'), (int, float)) else f"**Progress vs previous attempt:** {pr.get('overall', '')}"]
            L += [f"- Improved: {x}" for x in pr.get("improved", [])] + [f"- Not improved: {x}" for x in pr.get("not_improved", [])]
        L += ["", "**Model answer:**", "", ev.get("model_answer", "")]
    with open(os.path.join(SESSIONS_DIR, sess["id"], "report.md"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))


# =====================================================================
#  SHADOW PRACTICE  (voice only: listen to a phrase, repeat it, get rated)
# =====================================================================
SHADOW_DIR = os.path.join(DATA_DIR, "shadow_sets")
SHADOW_PROGRESS_DIR = os.path.join(DATA_DIR, "shadow_progress")
TTS_CACHE_DIR = os.path.join(DATA_DIR, "tts_cache")
SHADOW_REC_DIR = os.path.join(BASE_DIR, "shadow_recordings")
SEED_DIR = os.path.join(BASE_DIR, "seed")
for d in (SHADOW_DIR, SHADOW_PROGRESS_DIR, TTS_CACHE_DIR, SHADOW_REC_DIR):
    os.makedirs(d, exist_ok=True)

def install_seed_sets():
    """Copy bundled starter sets (seed/shadow_*.json) into data/shadow_sets if they're not there yet."""
    import shutil
    if not os.path.isdir(SEED_DIR):
        return
    for fn in os.listdir(SEED_DIR):
        if fn.startswith("shadow_") and fn.endswith(".json"):
            dst = os.path.join(SHADOW_DIR, fn[len("shadow_"):])
            if not os.path.exists(dst):
                shutil.copy(os.path.join(SEED_DIR, fn), dst)


install_seed_sets()

ACCENTS = {
    "american": "Speak with a natural, neutral American English accent, like an experienced software engineer from the US in a job interview.",
    "indian": "Speak with a natural, clear Indian English accent, like an experienced software engineer from India speaking confidently in a job interview.",
}
STYLES = {
    "natural": ("Sound like an experienced engineer thinking while they talk in a real interview, NOT reading a script: "
                "relaxed and conversational, slightly uneven pace - a little slower on the key point, quicker on familiar parts, "
                "stress the most important word in each sentence, let your pitch fall at the end of statements, "
                "say fillers like 'um' or 'so' lightly and naturally, and treat '...' as a short thinking pause."),
    "plain": "",
}
PACES = {
    "normal": "Use a calm, confident, conversational pace with natural pauses between ideas.",
    "slow": "Speak slowly and very clearly, pausing slightly between phrases, so a learner can repeat after you.",
}


def split_phrases(text, max_words=16):
    """Split an answer into short phrases that are easy to repeat."""
    text = re.sub(r"\s+", " ", text or "").strip()
    sentences = re.split(r"(?<=[.!?])\s+(?=[A-Z0-9\"'])", text)
    out = []
    for s in sentences:
        s = s.strip()
        if not s:
            continue
        words = s.split()
        if len(words) <= max_words:
            out.append(s)
            continue
        # split long sentences at commas / conjunctions near the middle
        parts, cur = [], []
        for w in words:
            cur.append(w)
            brk = w.endswith((",", ";", ":")) or (w.lower() in ("and", "but", "so", "because", "which", "then") and len(cur) > 1)
            if brk and len(cur) >= 6:
                if w.lower() in ("and", "but", "so", "because", "which", "then"):
                    cur.pop()
                    parts.append(" ".join(cur))
                    cur = [w]
                else:
                    parts.append(" ".join(cur))
                    cur = []
            elif len(cur) >= max_words + 4:
                parts.append(" ".join(cur))
                cur = []
        if cur:
            if parts and len(cur) < 4:
                parts[-1] += " " + " ".join(cur)
            else:
                parts.append(" ".join(cur))
        out.extend(parts)
    # merge very short bits ("Sure.", "Second,") into the following phrase
    merged = []
    for p in out:
        if merged and len(merged[-1].split()) < 4:
            merged[-1] += " " + p
        else:
            merged.append(p)
    if len(merged) > 1 and len(merged[-1].split()) < 3:
        merged[-2] += " " + merged.pop()
    return merged


def shadow_path(set_id):
    return os.path.join(SHADOW_DIR, slug(set_id, 80) + ".json")


def load_shadow(set_id):
    return read_json(shadow_path(set_id))


def save_shadow(s):
    write_json(shadow_path(s["id"]), s)


def progress_path(set_id):
    return os.path.join(SHADOW_PROGRESS_DIR, slug(set_id, 80) + ".json")


def read_file_text(f):
    name = (f.filename or "").lower()
    data = f.read()
    if name.endswith(".pdf"):
        return parse_pdf(data)
    if name.endswith(".docx"):
        return parse_docx(data)
    return data.decode("utf-8", errors="ignore")


ANSWER_SYSTEM = """You write interview answers that the candidate will practice OUT LOUD by shadowing (listening and repeating).
Write exactly what the candidate would SAY, in first person, as an experienced professional in a real interview.
Rules:
- Use ONLY experience, companies, projects and skills that appear in the candidate's resume. Never invent employers,
  numbers, percentages, team sizes or metrics. If the question needs experience the resume doesn't show, answer from
  solid technical knowledge ("The way I'd approach it is...") without claiming fake experience.
- Natural spoken English: short sentences (mostly under 18 words), contractions, clear signposting
  ("First...", "The key thing is...", "For example, in my last project..."), no bullet points, no markdown, no headings.
- 120-170 words (about 60-75 seconds spoken). Behavioral questions: STAR structure in natural speech.
- For 'questions to ask the interviewer', write how the candidate would naturally ask it, with one short line on why they're asking.
Return only the answer text."""


@app.get("/api/shadow/sets")
def shadow_sets():
    install_seed_sets()
    out = []
    for fn in sorted(os.listdir(SHADOW_DIR)):
        if fn.endswith(".json"):
            s = read_json(os.path.join(SHADOW_DIR, fn), {})
            items = s.get("items", [])
            out.append({"id": s.get("id"), "name": s.get("name"), "count": len(items),
                        "with_answers": sum(1 for i in items if i.get("answer"))})
    return jsonify(out)


@app.get("/api/shadow/sets/<set_id>")
def shadow_set(set_id):
    s = load_shadow(set_id)
    if not s:
        return jsonify({"error": "not found"}), 404
    prog = read_json(progress_path(set_id), {}) or {}
    for it in s["items"]:
        it["phrases"] = split_phrases(it.get("answer", ""))
        runs = prog.get(str(it["id"]), [])
        it["runs"] = len(runs)
        it["best"] = max([r.get("score", 0) for r in runs], default=None)
        it["last"] = runs[-1].get("score") if runs else None
    s.pop("resume_text", None)
    s["has_resume"] = bool(load_shadow(set_id).get("resume_text"))
    return jsonify(s)


@app.post("/api/shadow/sets")
def shadow_create():
    name = request.form.get("name") or "Shadow set"
    qs = []
    if "questions_file" in request.files and request.files["questions_file"].filename:
        f = request.files["questions_file"]
        n = f.filename.lower()
        data = f.read()
        if n.endswith((".xlsx", ".xlsm")):
            qs = parse_excel(data)
        elif n.endswith(".csv"):
            qs = parse_csv(data)
        else:
            text = parse_pdf(data) if n.endswith(".pdf") else parse_docx(data) if n.endswith(".docx") else data.decode("utf-8", "ignore")
            qs = split_questions_from_text(text)
    if request.form.get("questions_text"):
        qs += split_questions_from_text(request.form["questions_text"])
    if not qs:
        return jsonify({"error": "No questions found. Upload a file or paste questions."}), 400
    resume = ""
    if "resume_file" in request.files and request.files["resume_file"].filename:
        resume = read_file_text(request.files["resume_file"])
    resume = (resume + "\n" + (request.form.get("resume_text") or "")).strip()
    sid = slug(name, 60) + "_" + datetime.now().strftime("%m%d%H%M%S")
    s = {"id": sid, "name": name, "resume_text": resume, "job_description": request.form.get("job_description", ""),
         "items": [{"id": i + 1, "n": i + 1, "topic": q.get("category", ""), "high": False,
                    "question": q["question"], "answer": q.get("reference_answer", "")} for i, q in enumerate(qs)]}
    save_shadow(s)
    return jsonify({"id": sid})


@app.delete("/api/shadow/sets/<set_id>")
def shadow_delete(set_id):
    for p in (shadow_path(set_id), progress_path(set_id)):
        if os.path.exists(p):
            os.remove(p)
    return jsonify({"ok": True})


@app.post("/api/shadow/sets/<set_id>/items/<int:item_id>")
def shadow_update_item(set_id, item_id):
    body = request.get_json(force=True) or {}
    with LOCK:
        s = load_shadow(set_id)
        if not s:
            return jsonify({"error": "not found"}), 404
        it = next((i for i in s["items"] if i["id"] == item_id), None)
        if not it:
            return jsonify({"error": "not found"}), 404
        if "answer" in body:
            it["answer"] = (body["answer"] or "").strip()
        save_shadow(s)
    return jsonify({"ok": True, "phrases": split_phrases(it["answer"])})


@app.post("/api/shadow/sets/<set_id>/items/<int:item_id>/generate")
def shadow_generate(set_id, item_id):
    s = load_shadow(set_id)
    if not s:
        return jsonify({"error": "not found"}), 404
    it = next((i for i in s["items"] if i["id"] == item_id), None)
    if not it:
        return jsonify({"error": "not found"}), 404
    st = load_settings()
    user = (f"Candidate resume:\n{(s.get('resume_text') or '(no resume provided)')[:12000]}\n\n"
            f"Target role: {st.get('target_role')}\n"
            f"Job description: {(s.get('job_description') or st.get('job_description') or 'not provided')[:4000]}\n\n"
            f"Topic: {it.get('topic') or '-'}\nInterview question: {it['question']}")
    try:
        ans = chat_text(ANSWER_SYSTEM, user, max_tokens=1200).strip()
    except Exception as e:
        return jsonify({"error": str(e)}), 400
    ans = re.sub(r"^\s*(answer\s*:)\s*", "", ans, flags=re.I).strip().strip('"')
    with LOCK:
        s = load_shadow(set_id)
        for i in s["items"]:
            if i["id"] == item_id:
                i["answer"] = ans
        save_shadow(s)
    return jsonify({"answer": ans, "phrases": split_phrases(ans)})


# ------------------------------------------------------------------ TTS
@app.post("/api/tts")
def tts():
    import hashlib
    body = request.get_json(force=True) or {}
    text = (body.get("text") or "").strip()[:3000]
    if not text:
        return jsonify({"error": "empty"}), 400
    accent = body.get("accent") if body.get("accent") in ACCENTS else "american"
    pace = body.get("pace") if body.get("pace") in PACES else "normal"
    voice = re.sub(r"[^a-z]", "", (body.get("voice") or "alloy").lower()) or "alloy"
    style = body.get("style") if body.get("style") in STYLES else "plain"
    model = load_settings().get("tts_model") or "gpt-4o-mini-tts"
    key = hashlib.sha1((f"{model}|{voice}|{accent}|{pace}|{text}" + (f"|{style}" if style != "plain" else "")).encode()).hexdigest()
    fn = key + ".mp3"
    path = os.path.join(TTS_CACHE_DIR, fn)
    if not os.path.exists(path):
        try:
            kwargs = dict(model=model, voice=voice, input=text, response_format="mp3")
            if not model.startswith("tts-"):
                kwargs["instructions"] = (ACCENTS[accent] + " " + PACES[pace] + " " + STYLES[style]).strip()
            elif pace == "slow":
                kwargs["speed"] = 0.85
            resp = client().audio.speech.create(**kwargs)
            data = resp.read() if hasattr(resp, "read") else resp.content
            with open(path + ".tmp", "wb") as fh:
                fh.write(data)
            os.replace(path + ".tmp", path)
        except Exception as e:
            return jsonify({"error": f"OpenAI voice failed: {e}"}), 400
    return jsonify({"url": f"/tts/{fn}"})


@app.route("/tts/<fn>")
def tts_file(fn):
    if not re.fullmatch(r"[0-9a-f]{40}\.mp3", fn):
        abort(404)
    return send_from_directory(TTS_CACHE_DIR, fn, mimetype="audio/mpeg")


# ------------------------------------------------------------- scoring
CONTRACTIONS = {"i'm": "i am", "it's": "it is", "that's": "that is", "don't": "do not", "doesn't": "does not",
                "can't": "cannot", "won't": "will not", "isn't": "is not", "aren't": "are not", "i've": "i have",
                "i'd": "i would", "i'll": "i will", "we're": "we are", "we've": "we have", "we'd": "we would",
                "they're": "they are", "there's": "there is", "what's": "what is", "let's": "let us",
                "you're": "you are", "didn't": "did not", "wasn't": "was not", "shouldn't": "should not",
                "couldn't": "could not", "here's": "here is", "who's": "who is", "you'd": "you would"}
NUMS = {"0": "zero", "1": "one", "2": "two", "3": "three", "4": "four", "5": "five", "6": "six", "7": "seven",
        "8": "eight", "9": "nine", "10": "ten", "20": "twenty", "50": "fifty", "90": "ninety", "100": "hundred"}


def norm_words(text):
    t = (text or "").lower().replace("’", "'").replace("-", " ").replace("/", " ")
    out = []
    for w in re.findall(r"[a-z0-9']+", t):
        w = w.strip("'")
        if not w:
            continue
        w = CONTRACTIONS.get(w, w)
        w = NUMS.get(w, w)
        out.extend(w.split())
    return out


def score_repetition(target, heard, user_sec, ref_sec, vm=None):
    import difflib
    fillers = {"um", "uh", "erm", "hmm", "ah", "uhm"}
    tw = [w for w in norm_words(target) if w not in fillers]
    hw = [w for w in norm_words(heard) if w not in fillers]
    sm = difflib.SequenceMatcher(a=tw, b=hw, autojunk=False)
    words, matched, missed, wrong = [], 0, [], []
    for op, a1, a2, b1, b2 in sm.get_opcodes():
        if op == "equal":
            matched += a2 - a1
            words += [{"w": w, "s": "ok"} for w in tw[a1:a2]]
        elif op == "replace":
            ta, hb = tw[a1:a2], hw[b1:b2]
            if "".join(ta) == "".join(hb) or difflib.SequenceMatcher(a="".join(ta), b="".join(hb)).ratio() >= 0.85:
                matched += len(ta)
                words += [{"w": w, "s": "ok"} for w in ta]
            else:
                for i, w in enumerate(ta):
                    h = hb[i] if i < len(hb) else ""
                    if h and difflib.SequenceMatcher(a=w, b=h).ratio() >= 0.8:
                        matched += 1
                        words.append({"w": w, "s": "ok"})
                    else:
                        words.append({"w": w, "s": "wrong", "heard": h})
                        wrong.append({"said": h, "expected": w})
        elif op == "delete":
            words += [{"w": w, "s": "missed"} for w in tw[a1:a2]]
            missed += tw[a1:a2]
    soft = {"so", "and", "well", "like", "okay", "right", "yeah", "you", "know"}
    extra = sum(1 for op, a1, a2, b1, b2 in sm.get_opcodes() if op == "insert" for w in hw[b1:b2] if w not in soft)
    accuracy = round(100 * matched / max(len(tw), 1))
    # pace: ratio of your speaking time vs the model voice
    pace_ratio = round(user_sec / ref_sec, 2) if (user_sec and ref_sec) else None
    if pace_ratio is None:
        pace_score = 80
    elif 0.8 <= pace_ratio <= 1.3:
        pace_score = 100
    else:
        pace_score = max(0, round(100 - abs((pace_ratio - (1.3 if pace_ratio > 1.3 else 0.8))) * 120))
    # fluency: long hesitations inside the phrase
    fluency = 100
    if vm:
        fluency -= 15 * (vm.get("pauses_over_0_7s") or 0) + 15 * (vm.get("pauses_over_2s") or 0)
        fluency = max(0, fluency)
    extra_pen = min(15, extra * 3)
    score = max(0, round(0.7 * accuracy + 0.15 * pace_score + 0.15 * fluency - extra_pen))
    return {"score": score, "accuracy": accuracy, "pace_ratio": pace_ratio, "pace_score": pace_score,
            "fluency": fluency, "extra_words": extra, "words": words, "missed": missed, "wrong": wrong,
            "target_words": len(tw), "heard_words": len(hw)}


@app.post("/api/shadow/score")
def shadow_score():
    set_id = request.form.get("set_id", "")
    item_id = request.form.get("item_id", "0")
    pidx = request.form.get("phrase_index", "0")
    target = request.form.get("target", "")
    a = request.files.get("audio")
    if not a:
        return jsonify({"error": "no audio"}), 400
    folder = os.path.join(SHADOW_REC_DIR, slug(set_id, 80), f"q{int(item_id):03d}")
    os.makedirs(folder, exist_ok=True)
    ext = ".mp4" if "mp4" in (a.mimetype or "") else ".ogg" if "ogg" in (a.mimetype or "") else ".webm"
    kind = re.sub(r"[^a-z-]", "", request.form.get("kind", "")) or "recall"
    fn = datetime.now().strftime("%Y%m%d_%H%M%S_") + (f"p{int(pidx):02d}" if pidx.lstrip("-").isdigit() and int(pidx) >= 0 else kind) + ext
    path = os.path.join(folder, fn)
    a.save(path)
    try:
        vm = json.loads(request.form.get("voice_metrics") or "null")
    except Exception:
        vm = None
    st = load_settings()
    try:
        oc, stt_model = stt()
        with open(path, "rb") as fh:
            tr = oc.audio.transcriptions.create(
                model=stt_model, file=fh,
                prompt="Spoken repetition of an interview answer. Technical terms: " + target[:600])
        heard = getattr(tr, "text", "") or ""
    except Exception as e:
        return jsonify({"error": f"Transcription failed: {e}"}), 400
    if vm and pidx.lstrip("-").isdigit() and int(pidx) < 0:
        vm = {"pauses_over_0_7s": 0, "pauses_over_2s": vm.get("pauses_over_2s", 0)}  # natural thinking pauses are fine
    res = score_repetition(target, heard, float(request.form.get("duration") or 0),
                           float(request.form.get("ref_duration") or 0), vm)
    res["heard"] = heard
    res["recording"] = f"/shadowrec/{slug(set_id, 80)}/q{int(item_id):03d}/{fn}"
    return jsonify(res)


@app.route("/shadowrec/<set_id>/<qdir>/<fn>")
def shadow_rec_file(set_id, qdir, fn):
    return send_from_directory(os.path.join(SHADOW_REC_DIR, slug(set_id, 80), slug(qdir, 10)), fn)


@app.post("/api/shadow/progress")
def shadow_progress():
    b = request.get_json(force=True) or {}
    p = progress_path(b.get("set_id", ""))
    with LOCK:
        prog = read_json(p, {}) or {}
        prog.setdefault(str(b.get("item_id")), []).append({
            "date": datetime.now().isoformat(timespec="seconds"), "score": b.get("score"),
            "mode": b.get("mode"), "accent": b.get("accent"), "phrases": b.get("phrases")})
        write_json(p, prog)
    return jsonify({"ok": True})


SHADOW_COACH_SYSTEM = """You are a spoken-English and interview delivery coach. The candidate practiced an interview answer
by listening to a model voice phrase by phrase and repeating it. You get, for each phrase, the target text, what the
speech recognizer heard, the words they missed or said differently, and pace/pause data. Words that were heard wrongly
are usually mispronounced or swallowed. 'rhythm' (if present) has silence_pct (natural speech: 15-25%), thinking_pauses
(pauses >= 0.7 s; natural answers have several at idea changes) and longest_pause_sec - even, short pauses with very
little silence sound read or recited; comment on this. Return ONLY JSON:
{"summary":"2-3 sentences, honest and encouraging",
 "pronunciation":[{"word":"...","tip":"how to say it: syllables / stress, e.g. 'em-BED-ding'"}],
 "fluency_and_pace":["..."],
 "sound_more_natural":["tips on rhythm, stressing key words, pausing between ideas, confident tone"],
 "next_step":"what to practice next"}"""


@app.post("/api/shadow/coach")
def shadow_coach():
    b = request.get_json(force=True) or {}
    try:
        res = chat_json(SHADOW_COACH_SYSTEM, json.dumps({"accent_practiced": b.get("accent"),
                                                         "results": b.get("results", [])[:40]}), max_tokens=1500)
        return jsonify(res)
    except Exception as e:
        return jsonify({"error": str(e)}), 400


# =====================================================================
#  AI LISTENS: delivery analysis by an audio-capable model
# =====================================================================
DELIVERY_DIR = os.path.join(DATA_DIR, "delivery_analyses")
os.makedirs(DELIVERY_DIR, exist_ok=True)

DELIVERY_PROMPT = """You are an expert interview and public-speaking coach with a trained ear. LISTEN to the attached recording
of a job candidate and judge HOW it sounds to a real interviewer - not just the words.

Context:
- What the candidate was doing: {mode}
- Interview question: {question}
- Script / content they were aiming for: {target}
- Accent they are practising: {accent}
- App's rough microphone measurements (approximate): {rhythm}

Listen for: whether it sounds spontaneous or recited/read (even rhythm, no thinking pauses, sing-song or flat pitch,
reading cadence), tone and warmth, confidence, pace changes, where they pause (at commas/full stops/idea changes vs
mid-phrase), sentence-final intonation (falling = confident, rising/flat = unsure or reading), emphasis on key words,
energy across the answer (does it fade at the end?), clarity, mumbling, swallowed endings and mispronounced words
(especially technical terms). Real conversation has uneven pauses: short breaths between words, longer 0.7-1.5 s pauses
when moving to a new idea, and roughly 15-25% silence. Be honest and specific; give approximate timestamps (m:ss).

Return ONLY a JSON object, no markdown:
{{
 "naturalness_score": number 0-10 (10 = sounds completely spontaneous and experienced),
 "sounds_like": "Natural & spontaneous | Prepared but natural | Rehearsed / recited | Read aloud",
 "would_interviewer_notice": "one honest sentence",
 "listener_impression": "2-3 sentences: how a hiring manager would perceive this person from the voice alone",
 "scores": {{"tone": n, "confidence": n, "pace": n, "pauses_and_phrasing": n, "intonation": n, "emphasis": n, "energy": n, "clarity_pronunciation": n}},
 "timeline": [{{"time": "0:00-0:12", "what_happens": "...", "rating": "good|ok|needs work", "fix": "..."}}],
 "strengths": ["..."],
 "top_fixes": ["the 3 most important changes, most important first"],
 "pauses_and_phrasing": ["specific observations about pauses, punctuation, run-on sentences, mid-phrase breaks"],
 "tone_and_intonation": ["specific observations"],
 "pronunciation": [{{"word": "...", "issue": "what it sounded like", "how_to_say": "e.g. LANG-graf, stress on first syllable"}}],
 "practice_plan": ["concrete step-by-step drills for the next practice session"]
}}"""


def _delivery_models(first):
    out = []
    for m in [first, "gpt-audio-mini", "gpt-audio", "gpt-4o-audio-preview", "gpt-4o-mini-audio-preview"]:
        if m and m not in out:
            out.append(m)
    return out


@app.post("/api/analyze_delivery")
def analyze_delivery():
    import hashlib
    b = request.get_json(force=True) or {}
    wav = b.get("wav_b64") or ""
    if not wav:
        return jsonify({"error": "No audio received."}), 400
    key = hashlib.sha1((b.get("recording") or wav[:5000]).encode()).hexdigest()
    cache = os.path.join(DELIVERY_DIR, key + ".json")
    cached = read_json(cache)
    if cached and not b.get("refresh"):
        cached["cached"] = True
        return jsonify(cached)
    st = load_settings()
    prompt = DELIVERY_PROMPT.format(
        mode=b.get("mode") or "answering an interview question",
        question=(b.get("question") or "unknown")[:500],
        target=(b.get("target") or "not provided")[:4000],
        accent=b.get("accent") or "not specified",
        rhythm=json.dumps(b.get("rhythm")) if b.get("rhythm") else "not available")
    last_err = None
    for model in _delivery_models(st.get("audio_model") or "gpt-audio-mini"):
        try:
            resp = client().chat.completions.create(
                model=model,
                modalities=["text"],
                messages=[{"role": "user", "content": [
                    {"type": "text", "text": prompt},
                    {"type": "input_audio", "input_audio": {"data": wav, "format": "wav"}}]}],
                max_completion_tokens=3000,
            )
            txt = resp.choices[0].message.content or ""
            m = re.search(r"\{.*\}", txt, re.S)
            res = json.loads(m.group(0)) if m else {"listener_impression": txt}
            res["model"] = model
            write_json(cache, res)
            return jsonify(res)
        except Exception as e:
            last_err = e
            msg = str(e).lower()
            # try the next model only when this one isn't available to the key
            if not any(w in msg for w in ("model", "not found", "does not exist", "access", "unsupported", "permission")):
                break
    return jsonify({"error": f"AI listening failed: {last_err}. Check the audio model in Settings."}), 400


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    if not os.environ.get("NO_BROWSER"):
        threading.Timer(1.2, lambda: webbrowser.open(f"http://127.0.0.1:{port}")).start()
    print(f"\n  AI Interviewer running at http://127.0.0.1:{port}\n  Recordings are saved in: {SESSIONS_DIR}\n")
    app.run(host="127.0.0.1", port=port, debug=False, threaded=True)
